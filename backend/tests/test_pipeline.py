import sys
import os
import pytest
import uuid
import time
import io
import openpyxl
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.main import app
from app.core.database import SessionLocal
from app.models.domain import Business, Campaign, PipelineJob, ExportHistory, BusinessContact, BusinessLocation
from app.services.pipeline import start_pipeline_job, get_pipeline_job_status, cancel_pipeline_job, resume_pipeline_job
from app.services.export_excel import generate_campaign_leads_xlsx
from app.services.export import generate_campaign_leads_csv

client = TestClient(app)


def test_pipeline_job_orchestration():
    """Test full pipeline job lifecycle: start, poll, completed status."""
    db = SessionLocal()
    try:
        campaign = Campaign(
            name="Pipeline Test Campaign",
            category="restaurants",
            locations=["Mumbai"],
            target_leads=5,
            provider="mock"
        )
        db.add(campaign)
        db.commit()
        camp_id = str(campaign.id)
    finally:
        db.close()

    res = start_pipeline_job(camp_id, batch_size=2)
    assert res["status"] == "success"
    job_id = res["job_id"]
    assert job_id is not None

    # Wait up to 5 seconds for worker completion
    completed = False
    for _ in range(25):
        time.sleep(0.2)
        st = get_pipeline_job_status(job_id)
        if st.get("job_status") in ["COMPLETED", "FAILED"]:
            completed = True
            assert st["job_status"] == "COMPLETED"
            assert st["progress_percent"] == 100.0
            break

    assert completed is True


def test_pipeline_cancellation_and_resumption():
    """Test cancelling and resuming a pipeline job."""
    db = SessionLocal()
    try:
        campaign = Campaign(
            name="Cancel Test Campaign",
            category="dentists",
            locations=["Delhi"],
            target_leads=10,
            provider="mock"
        )
        db.add(campaign)
        db.commit()
        camp_id = str(campaign.id)
    finally:
        db.close()

    res = start_pipeline_job(camp_id, batch_size=2)
    job_id = res["job_id"]

    # Cancel job
    cancel_res = cancel_pipeline_job(job_id)
    assert cancel_res["status"] == "success"

    st_cancelled = get_pipeline_job_status(job_id)
    assert st_cancelled["job_status"] == "CANCELLED"

    # Resume job
    resume_res = resume_pipeline_job(job_id)
    assert resume_res["status"] == "success"

    # Wait for completion after resume
    completed = False
    for _ in range(25):
        time.sleep(0.2)
        st = get_pipeline_job_status(job_id)
        if st.get("job_status") in ["COMPLETED", "FAILED"]:
            completed = True
            assert st["job_status"] == "COMPLETED"
            break

    assert completed is True


def test_excel_export_generation_multisheet():
    """Test generating 3-sheet Excel workbook (SALES LEADS, SUMMARY, DATA QUALITY)."""
    db = SessionLocal()
    try:
        campaign = Campaign(
            name="Excel Export Campaign",
            category="clinics",
            locations=["Chennai"]
        )
        db.add(campaign)
        db.commit()

        business = Business(
            name="Chennai Dental Care",
            category="clinics",
            website="https://chennaidental.com",
            rating=4.7,
            review_count=90
        )
        business.contacts = [BusinessContact(type="phone", value="+91 44 12345678")]
        business.locations = [BusinessLocation(address="Anna Salai", city="Chennai")]
        db.add(business)
        db.commit()

        xlsx_bytes = generate_campaign_leads_xlsx(campaign.id)
        assert len(xlsx_bytes) > 0

        # Parse Excel workbook using openpyxl from BytesIO stream
        wb = openpyxl.load_workbook(filename=io.BytesIO(xlsx_bytes), data_only=True)
        sheet_names = wb.sheetnames

        assert "SALES LEADS" in sheet_names
        assert "SUMMARY" in sheet_names
        assert "DATA QUALITY" in sheet_names

        ws1 = wb["SALES LEADS"]
        assert ws1.cell(row=1, column=1).value == "Lead ID"
        assert ws1.cell(row=1, column=2).value == "Business Name"

        # Check export history record
        hist = db.query(ExportHistory).filter(ExportHistory.campaign_id == campaign.id).first()
        assert hist is not None
        assert hist.format == "xlsx"
        assert hist.record_count >= 1
    finally:
        db.close()


def test_pipeline_api_endpoints():
    """Test REST API endpoints for pipeline and exports."""
    db = SessionLocal()
    try:
        campaign = Campaign(
            name="API Pipeline Campaign",
            category="bakeries",
            locations=["Pune"]
        )
        db.add(campaign)
        db.commit()
        camp_id = str(campaign.id)
    finally:
        db.close()

    # 1. Run Pipeline API
    resp_run = client.post(f"/api/v1/campaigns/{camp_id}/run-pipeline", json={"batch_size": 5})
    assert resp_run.status_code == 200
    job_data = resp_run.json()
    assert "job_id" in job_data
    job_id = job_data["job_id"]

    # 2. Get Job Status API
    resp_st = client.get(f"/api/v1/jobs/{job_id}")
    assert resp_st.status_code == 200
    st_data = resp_st.json()
    assert st_data["job_id"] == job_id
    assert "stage_progress" in st_data

    # 3. Export XLSX API
    resp_xlsx = client.post(f"/api/v1/campaigns/{camp_id}/export/xlsx", json={"qualified_only": False})
    assert resp_xlsx.status_code == 200
    assert "spreadsheetml.sheet" in resp_xlsx.headers["content-type"]

    # 4. Get Export History API
    resp_hist = client.get(f"/api/v1/campaigns/{camp_id}/exports")
    assert resp_hist.status_code == 200
    hist_list = resp_hist.json()
    assert len(hist_list) >= 1
